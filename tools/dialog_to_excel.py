#!/usr/bin/env python3
"""Gör om JSON-raderna från dialog_export.mjs till en redigerbar Excel-fil.

Användning:
    node tools/dialog_export.mjs > /tmp/dialog.json
    python3 tools/dialog_to_excel.py /tmp/dialog.json Dialog.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def main():
    if len(sys.argv) != 3:
        print("Användning: dialog_to_excel.py in.json out.xlsx", file=sys.stderr)
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as f:
        rows = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "Dialog"

    headers = ["Fil", "Sammanhang", "Fält", "Text", "Nyckel (rör ej)"]
    ws.append(headers)
    header_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"

    lock_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        ws.append([row["file"], row["breadcrumb"], row["field"], row["text"], row["key"]])
        r = ws.max_row
        ws.cell(row=r, column=1).fill = lock_fill
        ws.cell(row=r, column=2).fill = lock_fill
        ws.cell(row=r, column=3).fill = lock_fill
        ws.cell(row=r, column=4).alignment = wrap
        ws.cell(row=r, column=5).fill = lock_fill
        ws.cell(row=r, column=5).font = Font(color="999999", size=8)

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 90
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["E"].hidden = True

    # Skydda allt utom Text-kolumnen mot redigering (rader läggs inte till/tas bort
    # av misstag lika lätt då). Lösenord sätts inte - det är bara ett skydd mot
    # klantiga misstag, inte en säkerhetsspärr.
    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row_cells:
            cell.protection = cell.protection.copy(locked=cell.column_letter != "D")
    for cell in ws[1]:
        cell.protection = cell.protection.copy(locked=True)
    ws.protection.sheet = True
    ws.protection.formatPassword = ""
    ws.protection.enable()

    wb.save(sys.argv[2])
    print(f"Skrev {len(rows)} rader till {sys.argv[2]}")


if __name__ == "__main__":
    main()
