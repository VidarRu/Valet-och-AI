#!/usr/bin/env python3
"""Läser en redigerad Dialog.xlsx och gör om den till JSON som
tools/dialog_import.mjs kan skriva tillbaka till källkoden.

Användning:
    python3 tools/dialog_from_excel.py Dialog.xlsx | node tools/dialog_import.mjs
"""
import json
import sys

from openpyxl import load_workbook


def main():
    if len(sys.argv) != 2:
        print("Användning: dialog_from_excel.py in.xlsx", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(sys.argv[1])
    ws = wb.active

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        _file, _breadcrumb, _field, text, key = row
        if key is None:
            continue  # tom rad
        rows.append({"key": key, "text": text or ""})

    json.dump(rows, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main()
